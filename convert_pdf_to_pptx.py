import os
import sys
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

pdf_path = "slide_presentasi2.pdf"
pptx_path = "slide_presentasi2_visual.pptx"

if not os.path.exists(pdf_path):
    print(f"Error: {pdf_path} not found")
    sys.exit(1)

print("Converting PDF pages to images...")
# Convert PDF to images at 300 DPI for good quality
images = convert_from_path(pdf_path, dpi=300, fmt="png")
print(f"Total pages: {len(images)}")

# Create presentation with 16:9 aspect ratio
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Use blank layout (index 6 usually, but let's find it)
blank_layout = None
for layout in prs.slide_layouts:
    if layout.name == "Blank":
        blank_layout = layout
        break
if blank_layout is None:
    blank_layout = prs.slide_layouts[6]  # fallback

output_dir = "pdf_images"
os.makedirs(output_dir, exist_ok=True)

for i, image in enumerate(images, start=1):
    img_path = os.path.join(output_dir, f"slide_{i:03d}.png")
    image.save(img_path, "PNG")
    
    slide = prs.slides.add_slide(blank_layout)
    # Add picture to fill the slide
    left = Inches(0)
    top = Inches(0)
    pic = slide.shapes.add_picture(img_path, left, top, width=prs.slide_width, height=prs.slide_height)
    
    if i % 10 == 0:
        print(f"Processed {i}/{len(images)} slides...")

prs.save(pptx_path)
print(f"Saved: {pptx_path}")
